# -*- coding: utf-8 -*-
"""Verify the built deck: slide count, leftover placeholders, out-of-bounds text."""
import sys, io
from pptx import Presentation
from pptx.util import Emu

path = sys.argv[1]
prs = Presentation(path)
W, H = prs.slide_width, prs.slide_height
bad_tokens = ["Lorem ipsum", "点击此处", "一级标题", "输入文字", "可替换图片",
              "添加标题文本", "编号", "稻壳儿", "原创来源", "建议字号"]
buf = io.StringIO()
def w(s=""):
    buf.write(s + "\n")

w("slides = %d  size = %.2f x %.2f in" % (len(prs.slides), W / 914400, H / 914400))
issues = 0
def iter_shapes(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            for c in iter_shapes(sh.shapes):
                yield c

for i, slide in enumerate(prs.slides, 1):
    texts = []
    for sh in iter_shapes(slide.shapes):
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if t:
                texts.append(t.replace("\n", " / "))
                for tok in bad_tokens:
                    if tok in t:
                        w("  !! slide %d has placeholder %r in %r" % (i, tok, t[:40]))
                        issues += 1
        if sh.has_table:
            for r in sh.table.rows:
                for c in r.cells:
                    if c.text.strip():
                        texts.append(c.text.strip())
    head = texts[0][:46] if texts else "(no text)"
    w("%02d  %s" % (i, head))
    # out of bounds (text boxes only)
    for sh in iter_shapes(slide.shapes):
        if not sh.has_text_frame:
            continue
        if sh.left is None or sh.width is None:
            continue
        if sh.left < -Emu(9144) or sh.top < -Emu(9144) or sh.left + sh.width > W + Emu(9144) or sh.top + sh.height > H + Emu(9144):
            w("     oob: %r (%.2f,%.2f %.2fx%.2f) text=%r" % (
                sh.name, sh.left / 914400, sh.top / 914400, sh.width / 914400, sh.height / 914400,
                sh.text_frame.text[:30]))
            issues += 1
w()
w("issues = %d" % issues)
open(sys.argv[2], "w", encoding="utf-8").write(buf.getvalue())
print("issues:", issues)
