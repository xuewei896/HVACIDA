# -*- coding: utf-8 -*-
"""Build the HVACIDA project presentation on top of the CREC template.

Structure: template cover / contents / section pages / end page are reused,
content pages are built on the template's '标题和内容' layout so that the
company footer, logo and accent band come from the template itself.
"""
import copy
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION

TEMPLATE = r"D:\DSH\PPT模板-中铁六院(中铁六院集团公司电气化设计院分公司).pptx"
OUT = r"D:\DSH\结题汇报-基于Revit二次开发的机电系统BIM正向设计研究.pptx"

NAVY = RGBColor(0x27, 0x36, 0x77)
BLUE = RGBColor(0x3B, 0x53, 0x9D)
MID = RGBColor(0x4F, 0x78, 0xC9)
LIGHT = RGBColor(0x94, 0xB9, 0xE5)
PALE = RGBColor(0xEE, 0xF3, 0xFA)
PALE2 = RGBColor(0xF7, 0xFA, 0xFD)
TEXT = RGBColor(0x33, 0x33, 0x33)
GRAY = RGBColor(0x6B, 0x7B, 0x8C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x21, 0x2A, 0x4D)
FONT = "微软雅黑"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"

TOTAL = 21


# ---------------------------------------------------------------- helpers ----
def iter_shapes(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            for child in iter_shapes(sh.shapes):
                yield child


def find(slide, name):
    for sh in iter_shapes(slide.shapes):
        if sh.name == name:
            return sh
    return None


def rewrite(shape, lines, size=12, bold=False, color=TEXT, font=FONT,
            align=None, line_spacing=None):
    """Replace all text of a shape with `lines` (same formatting for each)."""
    if shape is None or not shape.has_text_frame:
        return
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if align is not None:
            p.alignment = align
        if line_spacing is not None:
            p.line_spacing = line_spacing
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = font
        r.font.color.rgb = color


def textbox(slide, x, y, w, h, lines, size=12, bold=False, color=TEXT,
            align=PP_ALIGN.LEFT, line_spacing=1.22, anchor=MSO_ANCHOR.TOP,
            space_after=4, wrap=True, font=FONT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(lines, str):
        lines = [lines]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        if i < len(lines) - 1:
            p.space_after = Pt(space_after)
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = font
        r.font.color.rgb = color
    return tb


def bullets(slide, x, y, w, h, items, size=11, color=TEXT, gap=6,
            mark="▪ ", mark_color=None, line_spacing=1.28):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = line_spacing
        p.space_after = Pt(gap)
        if mark:
            rm = p.add_run()
            rm.text = mark
            rm.font.size = Pt(size)
            rm.font.name = FONT
            rm.font.color.rgb = mark_color or MID
        r = p.add_run()
        r.text = item
        r.font.size = Pt(size)
        r.font.name = FONT
        r.font.color.rgb = color
    return tb


def panel(slide, x, y, w, h, fill=PALE, line=None, radius=0.05, shadow=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                   Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.75)
    shape.shadow.inherit = shadow
    try:
        shape.adjustments[0] = radius
    except Exception:
        pass
    return shape


def kpi(slide, x, y, w, h, value, label, value_size=22):
    panel(slide, x, y, w, h, PALE, None, 0.08)
    bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x + 0.02), Inches(y + 0.14),
                                 Inches(0.055), Inches(h - 0.28))
    bar.fill.solid(); bar.fill.fore_color.rgb = MID
    bar.line.fill.background()
    textbox(slide, x + 0.20, y + 0.07, w - 0.28, 0.40, [value], size=value_size, bold=True,
            color=BLUE, line_spacing=1.0)
    textbox(slide, x + 0.20, y + 0.06 + value_size / 72.0 * 1.28, w - 0.26, 0.24, [label],
            size=8.8, color=GRAY, line_spacing=1.0)


def header(slide, title, label=None):
    textbox(slide, 0.70, 0.06, 8.55, 0.40, [title], size=19, bold=True, color=INK, line_spacing=1.0)
    if label:
        textbox(slide, 0.72, 0.42, 8.3, 0.24, [label], size=8, bold=False, color=GRAY, line_spacing=1.0)


def page_no(slide, n):
    textbox(slide, 8.75, 5.235, 0.95, 0.24, ["%02d / %d" % (n, TOTAL)], size=8.5, color=WHITE,
            align=PP_ALIGN.RIGHT, line_spacing=1.0)


def add_page(prs, title, label, n):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    header(slide, title, label)
    page_no(slide, n)
    return slide


def table(slide, x, y, w, rows, col_pct, font_size=9, row_h=0.30, head_h=0.32,
          head_fill=NAVY, align_first_left=True):
    nrows, ncols = len(rows), len(rows[0])
    height = head_h + row_h * (nrows - 1)
    shape = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(height))
    tbl = shape.table
    tbl.first_row = True
    tbl.horz_banding = False
    total = sum(col_pct)
    for ci, pct in enumerate(col_pct):
        tbl.columns[ci].width = Emu(int(Inches(w) * pct / total))
    tbl.rows[0].height = Inches(head_h)
    for ri in range(1, nrows):
        tbl.rows[ri].height = Inches(row_h)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.margin_left = Inches(0.05)
            cell.margin_right = Inches(0.05)
            cell.margin_top = Inches(0.015)
            cell.margin_bottom = Inches(0.015)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if (ci == 0 and align_first_left) else PP_ALIGN.CENTER
            p.line_spacing = 1.0
            r = p.add_run()
            r.text = str(val)
            r.font.name = FONT
            r.font.size = Pt(font_size)
            if ri == 0:
                r.font.bold = True
                r.font.color.rgb = WHITE
                cell.fill.solid(); cell.fill.fore_color.rgb = head_fill
            else:
                r.font.color.rgb = TEXT
                cell.fill.solid()
                cell.fill.fore_color.rgb = PALE2 if ri % 2 == 1 else WHITE
    return shape


def clone_slide(prs, src):
    """Duplicate a template slide (shapes + relationships + background)."""
    new = prs.slides.add_slide(src.slide_layout)
    for sh in src.shapes:
        new.shapes._spTree.append(copy.deepcopy(sh.element))
    for el in new.shapes._spTree.iter():
        for key in list(el.attrib):
            if not key.startswith(R):
                continue
            rel = src.part.rels.get(el.attrib[key])
            if rel is None:
                del el.attrib[key]
                continue
            if rel.is_external:
                el.attrib[key] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
            else:
                el.attrib[key] = new.part.relate_to(rel.target_part, rel.reltype)
    if src.element.cSld.bg is not None:
        new.element.cSld.insert(0, copy.deepcopy(src.element.cSld.bg))
    return new


def set_order(prs, ordered, drop):
    """Reorder slides to `ordered` and remove the parts of `drop`."""
    sldIdLst = prs.slides._sldIdLst
    ids = list(sldIdLst)
    slides = list(prs.slides)
    by_id = {}
    for sldId, slide in zip(ids, slides):
        by_id[id(slide)] = sldId
    for sldId in ids:
        sldIdLst.remove(sldId)
    for slide in ordered:
        sldIdLst.append(by_id[id(slide)])
    for slide in drop:
        rid = by_id[id(slide)].get(R + "id")
        if rid:
            try:
                prs.part.drop_rel(rid)
            except Exception:
                pass


# ---------------------------------------------------------------- deck ----
prs = Presentation(TEMPLATE)
tpl = list(prs.slides)

# ============================== 1. cover ==================================
s = tpl[0]
rewrite(find(s, "矩形 259"), ["基于 Revit 二次开发的", "机电系统 BIM 正向设计研究"],
        size=30, bold=True, color=WHITE, align=PP_ALIGN.LEFT, line_spacing=1.15)
eng = find(s, "矩形 16")
if eng is not None:
    eng._element.getparent().remove(eng._element)
textbox(s, 2.95, 3.18, 4.6, 0.22, ["BIM FORWARD DESIGN FOR MEP BASED ON REVIT API"],
        size=7.5, bold=False, color=RGBColor(0xDC, 0xE7, 0xF5), align=PP_ALIGN.LEFT, line_spacing=1.0)
lbl = find(s, "矩形 24")
if lbl is not None:
    lbl.width = Inches(2.30)
rewrite(lbl, ["汇报人: 韩学伟"], size=12, bold=True, color=WHITE, line_spacing=1.0)
lbl = find(s, "矩形 34")
if lbl is not None:
    lbl.width = Inches(2.30)
rewrite(lbl, ["联系方式: 机电设备所"], size=12, bold=True, color=WHITE, line_spacing=1.0)
lbl = find(s, "矩形 1")
if lbl is not None:
    lbl.width = Inches(1.70)
rewrite(lbl, ["日期: 2026.10"], size=12, bold=True, color=WHITE, line_spacing=1.0)
for sh in iter_shapes(s.shapes):
    if sh.has_text_frame and sh.text_frame.text.strip() == "编号":
        rewrite(sh, ["结题汇报"], size=10.65, bold=True, color=WHITE, align=PP_ALIGN.CENTER, line_spacing=1.0)
cover = s

# ============================== 2. contents ===============================
s = tpl[2]
toc_titles = ["01  研究背景与目标", "02  技术方案与实现", "03  验证与效果", "04  结论与展望"]
toc_desc = ["三个结构性问题、五项研究内容与技术路线",
            "计算内核、模型联动、成果派生、规范知识与 AI",
            "三套自动化门禁、算例复算与效率分析",
            "主要结论、创新点与后续工作建议"]
i = 0
for sh in iter_shapes(s.shapes):
    if sh.has_text_frame and sh.text_frame.text.strip() == "输入文字":
        rewrite(sh, [toc_titles[i]], size=17, bold=True, color=INK, align=PP_ALIGN.LEFT, line_spacing=1.0)
        i += 1
i = 0
for sh in iter_shapes(s.shapes):
    if sh.has_text_frame and "点击此处添加文本信息" in sh.text_frame.text:
        if i < 4:
            rewrite(sh, [toc_desc[i]], size=8, bold=False, color=GRAY, align=PP_ALIGN.LEFT, line_spacing=1.15)
        i += 1
toc = s

# ============================== sections ==================================
def section(src, number, title, desc):
    new = clone_slide(prs, src)
    for sh in iter_shapes(new.shapes):
        if not sh.has_text_frame:
            continue
        txt = sh.text_frame.text.strip()
        if txt in ("01", "02"):
            rewrite(sh, [number], size=54, bold=False, color=WHITE, align=PP_ALIGN.LEFT, line_spacing=1.0)
        elif "添加标题文本" in txt:
            rewrite(sh, [title], size=30, bold=True, color=WHITE, align=PP_ALIGN.LEFT, line_spacing=1.1)
        elif "点击此处添加文本信息" in txt:
            rewrite(sh, [desc], size=9.5, bold=True, color=WHITE, align=PP_ALIGN.LEFT, line_spacing=1.2)
    return new

sec1 = section(tpl[3], "01", "研究背景与目标", "地铁机电正向设计的三个断点与五项研究内容")
sec2 = section(tpl[3], "02", "技术方案与实现", "计算内核 · 模型联动 · 成果派生 · 规范知识与 AI")
sec3 = section(tpl[3], "03", "验证与效果", "三套自动化门禁与文档示例逐格复算")
sec4 = section(tpl[3], "04", "结论与展望", "主要结论、创新点与后续工作")

# ============================== 4. background =============================
s = add_page(prs, "机电正向设计的三个结构性问题", "BACKGROUND · WHY FORWARD DESIGN NEEDS A COMPUTING CORE", 4)
cards = [
    ("01", "口径漂移", "同一参数在计算表、说明、图纸里各取一值;\n改一处,其它处不同步,审核只能逐格比对。",
     "系数散落在个人表格中,无法追溯来源"),
    ("02", "重复录入", "模型里已有面积、层高,计算表再抄一遍;\n抄错一位小数,负荷与风量全部偏移。",
     "几何信息两个来源,错误不易被发现"),
    ("03", "成果割裂", "计算书、材料表、图纸清单各自维护;\n一次设计变更要改四处,统计口径未必一致。",
     "交付物之间缺少共同的源头"),
]
x = 0.55
for num, title, body, foot in cards:
    panel(s, x, 1.30, 2.90, 2.55, PALE, None, 0.06)
    textbox(s, x + 0.22, 1.44, 0.8, 0.45, [num], size=22, bold=True, color=LIGHT, line_spacing=1.0)
    textbox(s, x + 0.22, 1.92, 2.45, 0.32, [title], size=14, bold=True, color=INK, line_spacing=1.0)
    textbox(s, x + 0.22, 2.30, 2.48, 1.05, body.split("\n"), size=9.5, color=TEXT, line_spacing=1.35)
    textbox(s, x + 0.22, 3.40, 2.48, 0.36, [foot], size=8, color=GRAY, line_spacing=1.1)
    x += 3.03
panel(s, 0.55, 4.12, 8.90, 0.72, PALE2, MID, 0.10)
textbox(s, 0.80, 4.22, 8.4, 0.52,
        ["正向设计真正的要求不是「把图画进模型」,而是让 计算—建模—出图—算量 形成闭环:",
         "几何来自模型、负荷来自内核、交付物由同一份结果派生 —— 只有一个数据源。"],
        size=10.5, bold=True, color=BLUE, line_spacing=1.25)

# ============================== 5. objectives =============================
s = add_page(prs, "研究目标与五项研究内容", "OBJECTIVES · WHAT THE PROJECT DELIVERS", 5)
panel(s, 0.55, 1.28, 3.30, 3.30, NAVY, None, 0.05)
textbox(s, 0.80, 1.50, 2.85, 0.34, ["总目标"], size=13, bold=True, color=WHITE, line_spacing=1.0)
textbox(s, 0.80, 1.92, 2.85, 2.50,
        ["让设计计算在 BIM 模型上正向发生:口径唯一、结果可核、交付同源。",
         " ",
         "以地铁通风空调为对象,交付一套运行于 Revit 内的正向设计工具,并以自动化门禁保证口径不漂移。"],
        size=9.8, color=WHITE, line_spacing=1.4)
items = [
    ("计算内核化", "两份公式文档逐格迁移为纯计算代码,25 个算例场景、约 870 项断言"),
    ("模型取值", "公共区与房间几何从空间读取;两点测量解决模型给不出的长度"),
    ("成果派生", "计算书、Excel、材料表、图纸清单由同一份结果与同一份列定义生成"),
    ("知识与 AI", "口径、条文、操作指南一体化检索;大模型问答与受控工具调用"),
    ("质量体系", "数值 / 界面 / 结构三套门禁 + 57 条决策记录 + 文档同步"),
]
y = 1.30
for k, (name, desc) in enumerate(items, 1):
    panel(s, 4.10, y, 5.35, 0.58, PALE2, None, 0.12)
    textbox(s, 4.28, y + 0.11, 0.42, 0.32, ["%d" % k], size=15, bold=True, color=LIGHT, line_spacing=1.0)
    textbox(s, 4.74, y + 0.06, 1.35, 0.26, [name], size=10.5, bold=True, color=INK, line_spacing=1.0)
    textbox(s, 4.74, y + 0.29, 4.55, 0.24, [desc], size=8.5, color=GRAY, line_spacing=1.0)
    y += 0.64
textbox(s, 4.10, 4.50, 5.35, 0.44,
        ["课题:中铁六院电化院 2023 年院内创新课题 · 承担部门:机电设备所 · 负责人:韩学伟",
         "成果形式:研究报告 1 份、软件著作权 1 项、论文 1 篇"],
        size=8, bold=False, color=GRAY, line_spacing=1.35)
textbox(s, 4.10, 4.96, 5.35, 0.24,
        ["工程化:集中开发 47 天、105 次提交,每次提交前三套门禁必须全绿。"], size=8, bold=True,
        color=BLUE, line_spacing=1.0)

# ============================== 6. roadmap ===============================
s = add_page(prs, "技术路线:文档驱动 · 分层实现 · 门禁验证", "METHODOLOGY · FOUR STEPS, ONE DISCIPLINE", 6)
steps = [("1", "文档驱动", "公式文档为唯一权威\n保留单元格代号便于回溯\n示例用作逐格断言"),
         ("2", "分层实现", "Core 不引用 Revit\nUI 只引用 Core\nRevit 层只做适配"),
         ("3", "门禁验证", "数值 / 界面 / 结构\n离线可跑、无外部依赖\n失败即禁止提交"),
         ("4", "口径登记", "系数集中定义并注明来源\n冲突按文字实现\n57 条决策留痕")]
x = 0.55
for num, name, body in steps:
    panel(s, x, 1.30, 2.02, 2.20, PALE, None, 0.06)
    textbox(s, x + 0.18, 1.44, 1.7, 0.34, ["STEP " + num], size=8.5, bold=True, color=LIGHT, line_spacing=1.0)
    textbox(s, x + 0.18, 1.76, 1.8, 0.34, [name], size=13, bold=True, color=INK, line_spacing=1.0)
    textbox(s, x + 0.18, 2.18, 1.75, 1.20, body.split("\n"), size=8.8, color=TEXT, line_spacing=1.3)
    if x < 7.5:
        arrow = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + 2.08), Inches(2.28),
                                   Inches(0.28), Inches(0.22))
        arrow.fill.solid(); arrow.fill.fore_color.rgb = LIGHT
        arrow.line.fill.background()
    x += 2.42
panel(s, 0.55, 3.68, 4.30, 1.20, PALE2, None, 0.08)
textbox(s, 0.75, 3.80, 3.95, 0.26, ["冲突处理原则"], size=10.5, bold=True, color=BLUE, line_spacing=1.0)
textbox(s, 0.75, 4.08, 3.95, 0.72,
        ["文档文字与文档示例冲突时按文字实现;差异在界面与决策日志中写明,不擅自「凑数」。"],
        size=9, color=TEXT, line_spacing=1.3)
panel(s, 5.10, 3.68, 4.35, 1.20, PALE2, None, 0.08)
textbox(s, 5.30, 3.80, 4.0, 0.26, ["精度口径"], size=10.5, bold=True, color=BLUE, line_spacing=1.0)
textbox(s, 5.30, 4.08, 4.0, 0.72,
        ["几何参数统一保留 1 位小数且「显示位数即计算位数」;水力逐段尺寸按毫米建模,保持原精度。"],
        size=9, color=TEXT, line_spacing=1.3)

# ============================== 8. architecture ==========================
s = add_page(prs, "系统架构:三层单向依赖", "ARCHITECTURE · CORE / UI / REVIT", 8)
layers = [("HVACIDA.Revit", "24 文件 · 3 756 行", "应用入口与 Ribbon、命令、模型读写、拾取与两点测量、批量出图", NAVY),
          ("HVACIDA.UI", "41 + 17 文件 · 8 972 行", "16 个功能窗口 + 1 个复用结果控件、视图模型、转换器、列生成器", BLUE),
          ("HVACIDA.Core", "72 文件 · 12 974 行", "50 个服务:计算内核、结果表、知识库、持久化、XLSX 写入、AI 客户端", MID)]
y = 1.26
for name, size, desc, color in layers:
    panel(s, 0.55, y, 5.95, 1.02, color, None, 0.05)
    textbox(s, 0.78, y + 0.11, 2.05, 0.30, [name], size=12, bold=True, color=WHITE, line_spacing=1.0)
    textbox(s, 2.85, y + 0.14, 1.95, 0.26, [size], size=8.5, color=WHITE, line_spacing=1.0)
    textbox(s, 0.78, y + 0.46, 5.50, 0.50, [desc], size=9, color=WHITE, line_spacing=1.25)
    if y < 3.4:
        arrow = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(3.32), Inches(y + 1.04),
                                   Inches(0.22), Inches(0.16))
        arrow.fill.solid(); arrow.fill.fore_color.rgb = LIGHT
        arrow.line.fill.background()
    y += 1.22
panel(s, 6.75, 1.26, 2.75, 3.44, PALE, None, 0.06)
textbox(s, 6.95, 1.40, 2.35, 0.28, ["分层带来的可测试性"], size=10.5, bold=True, color=BLUE, line_spacing=1.0)
bullets(s, 6.95, 1.76, 2.38, 2.80,
        ["计算内核脱离 Revit 运行,构建机上可跑全部数值断言",
         "界面可在无模型环境下装载自检",
         "Revit 层只剩「取数」与「呈现」两类适配",
         "同一输入两次运行结果一致,无随机与时间依赖"],
        size=9, gap=8, line_spacing=1.3)
textbox(s, 0.55, 4.86, 5.95, 0.26, ["依赖方向:Revit → UI → Core,Core 不引用任何 Revit 程序集。"],
        size=8.5, bold=True, color=BLUE, line_spacing=1.0)

# ============================== 9. module map ============================
s = add_page(prs, "功能架构:7 个面板 · 24 个模块", "FUNCTION MAP · RIBBON AS SINGLE SOURCE OF TRUTH", 9)
rows = [["面板", "模块数", "主要功能", "关键输出"],
        ["项目信息", "2", "工程信息、气象参数(294 台站)", "project.xml · 气象联动"],
        ["大系统", "4", "公共区参数、负荷计算、排烟计算、计算结果", "分区负荷 / 风量 / 选型"],
        ["小系统", "7", "六类系统计算 + 全站汇总", "逐房间明细 · 设备选型"],
        ["水力计算", "3", "风系统、水系统、汇总", "总压损 · 扬程 · 不平衡率"],
        ["出图", "2", "材料表与图例表、图框与批量出图", "11 类构件统计 · 导出记录"],
        ["AI问答", "4", "操作指南、规范知识库、AI 对话、AI 助手", "检索增强问答 · 工具调用"],
        ["产品支持", "2", "问题反馈、帮助", "模块目录与口径说明"]]
table(s, 0.55, 1.30, 8.90, rows, [10, 8, 46, 36], font_size=9, row_h=0.36, head_h=0.36)
textbox(s, 0.55, 4.62, 8.90, 0.42,
        ["模块目录为单一数据源:Ribbon 按钮、图标、命令映射、帮助说明均由其派生,新增模块只需登记一次,门禁自动校验一致性。"],
        size=9, color=GRAY, line_spacing=1.25)

# ============================== 10. computing core =======================
s = add_page(prs, "计算内核:公式文档的可执行化", "COMPUTING CORE · CELL-BY-CELL VERIFICATION", 10)
kpi(s, 0.55, 1.28, 2.10, 0.94, "25", "个算例场景")
kpi(s, 2.82, 1.28, 2.10, 0.94, "870", "项断言")
kpi(s, 5.09, 1.28, 2.10, 0.94, "30", "项逐格校验")
kpi(s, 7.36, 1.28, 2.10, 0.94, "6", "类小系统")
rows = [["系统类型", "示例规模", "复算关键值", "结论"],
        ["全空气一次回风", "21 间 / 590.76 m²", "冷负荷 104.42 kW · 送风 32 539 m³/h", "一致"],
        ["多联机+新风", "11 间 / 187.01 m²", "21.57 kW · 湿负荷 1.733 g/s · 热湿比 12 448.6", "一致"],
        ["排风 / 排烟", "8 间 / 2 分区", "5 386 m³/h · 34 620 与 20 772 m³/h", "一致"],
        ["送风排风排烟", "5 间", "排风 20 307 · 送风 31 020 · 排烟 18 612", "一致"]]
table(s, 0.55, 2.42, 8.90, rows, [20, 20, 46, 14], font_size=9, row_h=0.32, head_h=0.32)
bullets(s, 0.55, 4.28, 8.90, 0.80,
        ["口径纪律:单元格代号只留在代码注释;交付文案不出现代号(门禁扫描);文档示例与文字冲突时按文字实现并留痕。",
         "边界保护:零面积、零人数、室内温度不高于过渡季温度、文件损坏、网络中断均有断言,不抛异常、不出负数。"],
        size=9, gap=4, line_spacing=1.25)

# ============================== 11. model link ===========================
s = add_page(prs, "模型信息获取与联动:让模型成为第一数据源", "MODEL LINK · SPACE, PICK, TWO-POINT MEASURE", 11)
rows = [["方式", "数据来源", "口径", "适用与局限"],
        ["空间属性读取", "空间面积、层高、体积", "保留 1 位小数", "房间与公共区几何;依赖空间已放置"],
        ["交互拾取", "多选空间快照", "同名去重,自动建行", "批量建房间表;需在模型中操作"],
        ["两点测量", "两次点选坐标求距", "英尺→米,保留 1 位", "外墙长度;需有效工作平面"],
        ["手工录入", "用户直接填写", "按录入精度", "无模型或模型不完整的场景"]]
table(s, 0.55, 1.28, 8.90, rows, [16, 24, 18, 42], font_size=9, row_h=0.34, head_h=0.34)
panel(s, 0.55, 3.34, 4.35, 1.55, PALE, None, 0.06)
textbox(s, 0.75, 3.46, 3.95, 0.26, ["开窗即同步(选项 A)"], size=10.5, bold=True, color=BLUE, line_spacing=1.0)
bullets(s, 0.75, 3.76, 3.95, 1.10,
        ["按来源空间标识重读面积 / 层高", "屋顶面积仅在等于旧面积时跟随",
         "空间被删 → 保留原值并点名,绝不写 0"], size=8.8, gap=3, line_spacing=1.2)
panel(s, 5.10, 3.34, 4.35, 1.55, PALE2, None, 0.06)
textbox(s, 5.30, 3.46, 3.95, 0.26, ["如实记录的边界"], size=10.5, bold=True, color=BLUE, line_spacing=1.0)
bullets(s, 5.30, 3.76, 3.95, 1.10,
        ["不做双向实时联动,模型改动需重开窗",
         "链接模型中的空间无法按标识重读",
         "外墙长度写入选中行,故选中行必须可靠"], size=8.8, gap=3, line_spacing=1.2)

# ============================== 12. deliverables =========================
s = add_page(prs, "成果派生:一份结果,多处一致", "DELIVERABLES · ONE SOURCE, MANY OUTPUTS", 12)
rows = [["交付物", "生成方式", "特点"],
        ["窗口结果表", "结果对象绑定,列定义来自 Core", "界面不自行拼装数据"],
        ["文本计算书", "结果格式化器固定版式", "含口径说明与待补项,可直接归档"],
        ["Excel 计算书", "自写最小 XLSX 写入器(零依赖)", "4 页:结果 / 明细 / 设备 / 口径"],
        ["材料表与图例表", "读模型 11 类构件后按类别规格归并", "归并键含单位,混合单位不相加"],
        ["图纸清单", "读图框统计 + 批量导出", "记录成功与失败原因"]]
table(s, 0.55, 1.28, 8.90, rows, [18, 42, 40], font_size=9, row_h=0.36, head_h=0.34)
panel(s, 0.55, 3.52, 8.90, 1.34, NAVY, None, 0.06)
textbox(s, 0.80, 3.64, 8.4, 0.28, ["同源带来的三个直接结果"], size=11, bold=True, color=WHITE, line_spacing=1.0)
textbox(s, 0.80, 3.98, 8.3, 0.80,
        ["① 一次输入变更,结果表、计算书、Excel 同时更新,不再分别维护;",
         "② 审核可按「结果—公式—系数—文档」回溯,不必猜测表格结构;",
         "③ 材料表与图例表共用同一份统计,永远不会互相矛盾。"],
        size=9.2, color=WHITE, line_spacing=1.35)

# ============================== 13. knowledge + AI =======================
s = add_page(prs, "规范知识与 AI 辅助:能力之外先讲边界", "KNOWLEDGE & AI · CONTROLLED BY DESIGN", 13)
blocks = [("三类知识一体化", "本项目口径 · 规范条文 · Revit 操作指南,同一检索入口;答不出时明说覆盖范围"),
          ("条文电子版导入", "txt / md / csv / docx 解析,识别条文号,幂等导入,中段关键词可搜"),
          ("检索增强问答", "先检索知识库再生成回答,提示词只带条文不带工程数据,错误码分级提示")]
y = 1.28
for name, desc in blocks:
    panel(s, 0.55, y, 4.35, 0.98, PALE, None, 0.08)
    textbox(s, 0.75, y + 0.12, 3.95, 0.26, [name], size=10.5, bold=True, color=INK, line_spacing=1.0)
    textbox(s, 0.75, y + 0.40, 3.95, 0.50, [desc], size=8.8, color=TEXT, line_spacing=1.25)
    y += 1.12
panel(s, 5.10, 1.28, 4.35, 3.24, NAVY, None, 0.05)
textbox(s, 5.32, 1.42, 3.9, 0.30, ["AI 的边界(六条)"], size=11, bold=True, color=WHITE, line_spacing=1.0)
bounds = ["两级开关,默认关闭", "命令白名单 + 逐条确认", "单次 ≤ 200 个构件、单事务",
          "数值参数一律不写回模型", "密钥加密存储,提示词不含工程数据",
          "边界写成门禁断言,「安全」可回归"]
yy = 1.78
for b in bounds:
    textbox(s, 5.32, yy, 3.95, 0.34, ["· " + b], size=9.5, color=WHITE, line_spacing=1.15)
    yy += 0.46
textbox(s, 0.55, 4.66, 4.35, 0.30, ["问答解决「查得到」,计算内核解决「算得对」。"],
        size=9, color=GRAY, line_spacing=1.0)

# ============================== 15. gates ================================
s = add_page(prs, "自动化门禁:三套检查,离线可跑", "QUALITY GATES · NO REVIT, NO NETWORK", 15)
rows = [["门禁", "检查内容", "规模", "失败含义"],
        ["核心算例", "25 个场景:算例回归、逐格复算、气象联动、水力、材料表、出图、知识库、AI 客户端", "约 870 项", "数值或业务规则不符"],
        ["窗口装载", "17 个界面构建、绑定完整性、按钮语义、参考表列数与可编辑性、转换器、选中行传递", "193 项", "界面不可用或行为回退"],
        ["Ribbon 结构", "面板与按钮数量、命令映射、图标存在性、模块目录一致性、空间类型写法扫描", "结构 + 扫描", "入口缺失或目录与实现不一致"]]
table(s, 0.55, 1.28, 8.90, rows, [14, 54, 14, 18], font_size=8.8, row_h=0.52, head_h=0.34)
bullets(s, 0.55, 3.42, 8.90, 1.40,
        ["可离线:控制台程序 + PowerShell 脚本,不需要 Revit 参与,也不需要人工点击;",
         "可读:失败输出断言名、期望值与实际值,直接指向问题;",
         "强制:提交前必须全绿,历史缺陷一旦复发会被立即拦住(如空间类型写法、交付文案出现单元格编号);",
         "回归:47 天 105 次提交,口径始终稳定,没有出现「越改越乱」。"],
        size=9.2, gap=3, line_spacing=1.3)

# ============================== 16. defect review ========================
s = add_page(prs, "典型缺陷复盘:根因,而不是表象", "LESSONS · SIX DEFECTS AND WHAT THEY TAUGHT", 16)
rows = [["现象", "根因", "处置与门禁"],
        ["打开参数窗报参数异常", "以 Space 直接做集合筛选,该类型不在原生对象模型", "改按 SpatialElement 收集再转换;静态扫描禁止回归"],
        ["删除的系统重开又出现", "保存只做增量更新,从不删除", "改为按类型整体替换,窗口状态权威"],
        ["添加行后界面不刷新", "参考表绑定合并行,只改了输入集合", "订阅集合变化并重建合并行"],
        ["测量结果总写入第一行", "参考表未绑定选中项,默认选中第一行", "合并行暴露底层对象,双向同步、忽略空值推送"],
        ["面积长度仍显示两位", "只改了结果表,录入列仍是裸绑定", "数值转换器 + 失焦回写,显示即计算精度"],
        ["合计行回风量恒为零", "该字段从未赋值", "按送风乘新风比折算,与全空气口径统一"]]
table(s, 0.55, 1.28, 8.90, rows, [24, 40, 36], font_size=8.8, row_h=0.44, head_h=0.34)
textbox(s, 0.55, 4.62, 8.90, 0.44,
        ["共同规律:缺陷几乎都出在「界面与领域对象的映射」与「保存语义」上;修复一律先定位根因,再补门禁断言,最后更新文档。"],
        size=9, color=GRAY, line_spacing=1.25)

# ============================== 17. project data =========================
s = add_page(prs, "项目数据:代码、门禁与决策的规模", "PROJECT DATA · CODE, TESTS, DECISIONS", 17)
kpi(s, 0.55, 1.28, 2.10, 0.94, "105", "次提交(集中开发期)")
kpi(s, 2.82, 1.28, 2.10, 0.94, "23 121", "行 C# 代码")
kpi(s, 5.09, 1.28, 2.10, 0.94, "24", "个功能模块")
kpi(s, 7.36, 1.28, 2.10, 0.94, "57", "条决策记录")
chart_data = CategoryChartData()
chart_data.categories = ["Core", "UI(含 XAML)", "Revit", "门禁脚本"]
chart_data.add_series("代码行数", (12974, 8972, 3756, 2212))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.55), Inches(2.40),
                        Inches(5.40), Inches(2.45), chart_data)
chart = gf.chart
chart.has_legend = False
chart.has_title = True
chart.chart_title.text_frame.text = "各层代码规模(行)"
for p in chart.chart_title.text_frame.paragraphs:
    for r in p.runs:
        r.font.size = Pt(9); r.font.name = FONT; r.font.bold = True
        r.font.color.rgb = INK
plot = chart.plots[0]
plot.gap_width = 70
plot.has_data_labels = True
dl = plot.data_labels
dl.font.size = Pt(8); dl.font.name = FONT
dl.position = XL_LABEL_POSITION.OUTSIDE_END
ser = plot.series[0]
ser.format.fill.solid(); ser.format.fill.fore_color.rgb = BLUE
va = chart.value_axis
va.visible = False
va.has_major_gridlines = False
ca = chart.category_axis
ca.tick_labels.font.size = Pt(8); ca.tick_labels.font.name = FONT
ca.format.line.color.rgb = RGBColor(0xC9, 0xD3, 0xE0)
panel(s, 6.15, 2.40, 3.30, 2.45, PALE, None, 0.06)
textbox(s, 6.35, 2.52, 2.95, 0.26, ["规模之外的三件事"], size=10.5, bold=True, color=BLUE, line_spacing=1.0)
bullets(s, 6.35, 2.84, 2.90, 1.85,
        ["约 1 100 项门禁断言,全部离线可跑",
         "48 个图标由内置工具批量生成",
         "7 份文档与决策日志同步维护"], size=9, gap=5, line_spacing=1.25)

# ============================== 19. conclusions ==========================
s = add_page(prs, "结论与创新点", "CONCLUSIONS · FOUR CONTRIBUTIONS", 19)
textbox(s, 0.55, 1.26, 4.30, 0.28, ["五条结论"], size=11, bold=True, color=BLUE, line_spacing=1.0)
conc = ["把公式文档迁移为可执行内核并逐格断言,是保证专业计算正确性的有效方法",
        "模型可作为几何第一来源,但要区分「模型给得出」「需交互指定」「模型给不出」三类数据",
        "成果派生必须同源,才能从结构上消除口径分叉",
        "AI 落地的关键不是模型能力,而是默认关闭、逐条确认、只读优先、可回滚的流程约束",
        "门禁 + 决策日志 + 文档同步,是一轮集中开发中保持口径稳定的可行实践"]
y = 1.62
for i, c in enumerate(conc, 1):
    textbox(s, 0.55, y, 0.34, 0.30, ["%d" % i], size=12, bold=True, color=LIGHT, line_spacing=1.0)
    textbox(s, 0.92, y - 0.02, 3.95, 0.62, [c], size=9, color=TEXT, line_spacing=1.28)
    lines = max(1, (len(c) + 20) // 21)
    y += 0.26 + 0.24 * lines
panel(s, 5.10, 1.26, 4.35, 3.62, PALE, None, 0.05)
textbox(s, 5.32, 1.40, 3.9, 0.28, ["四个创新点"], size=11, bold=True, color=BLUE, line_spacing=1.0)
innov = [("口径可执行、可追溯", "文档单元格—注释—界面提示—决策记录四者打通"),
         ("三套门禁的回归体系", "数值、界面、结构三层独立检查,均离线可跑"),
         ("拾取与编辑一致性闭环", "关窗—拾取—同视图模型重开;合并行双向选中"),
         ("面向工程可信的 AI 边界", "开关、白名单、数量上限、单事务、不写数值")]
yy = 1.76
for name, desc in innov:
    textbox(s, 5.32, yy, 3.95, 0.24, [name], size=9.5, bold=True, color=INK, line_spacing=1.0)
    textbox(s, 5.32, yy + 0.24, 3.95, 0.42, [desc], size=8.5, color=GRAY, line_spacing=1.2)
    yy += 0.70

# ============================== 20. limits / next ========================
s = add_page(prs, "局限与后续工作", "LIMITS & NEXT STEPS", 20)
rows = [["当前局限", "影响", "后续方向"],
        ["防烟分区几何未接入", "大系统排烟按整体取量,结果偏大", "接入分区面积、挡烟垂壁与储烟仓"],
        ["链接模型空间不能按标识重读", "链接模型下开窗同步不生效", "引入稳定标识映射,评估增量刷新"],
        ["同步为单向且需重开窗口", "模型改动不会即时反映", "事件驱动的增量同步"],
        ["消火栓水力与参数化构件未实装", "立项书所列消防水力、构件级信息联动未覆盖", "按假定流速法扩展至消防系统"],
        ["AI 不写数值、不批量改模型", "无法完成批量改参数类任务", "确认与回滚机制成熟后逐步放开"],
        ["仅在 Revit 2020 验证", "高版本 API 差异未覆盖", "收敛版本适配层,分版本门禁"],
        ["未实现 IFC 与多专业协同", "跨专业数据交换需另行处理", "评估 IFC 导出与模型检查"]]
table(s, 0.55, 1.24, 8.90, rows, [28, 34, 38], font_size=8.4, row_h=0.36, head_h=0.32)
panel(s, 0.55, 4.20, 8.90, 0.52, PALE2, MID, 0.10)
textbox(s, 0.80, 4.30, 8.4, 0.34,
        ["下一步最优先:防烟分区几何 + 稳定标识的模型联动 + 设备样本库,三者直接决定结果精度与使用体验。"],
        size=9.5, bold=True, color=BLUE, line_spacing=1.0)

# ============================== 21. end ==================================
s = tpl[17]
for sh in iter_shapes(s.shapes):
    if not sh.has_text_frame:
        continue
    txt = sh.text_frame.text
    if "感谢聆听" in txt:
        rewrite(sh, ["感谢聆听"], size=30, bold=False, color=WHITE, align=PP_ALIGN.CENTER, line_spacing=1.0)
    elif "Thank you" in txt:
        rewrite(sh, ["THANK YOU FOR LISTENING"], size=14, bold=True, color=WHITE,
                align=PP_ALIGN.CENTER, line_spacing=1.0)
    elif "汇报人" in txt:
        rewrite(sh, [""], size=8.25, bold=False, color=WHITE, align=PP_ALIGN.CENTER, line_spacing=1.2)
end = s
textbox(end, 1.25, 3.62, 2.95, 0.66,
        ["汇报人: HVACIDA 项目组", "日期: 2026 年 10 月"], size=9, bold=False, color=WHITE,
        align=PP_ALIGN.LEFT, line_spacing=1.5)

# ---------------------------------------------------------------- order ----
built = [sl for sl in prs.slides if all(sl is not t for t in tpl)]
sections = [sl for sl in built if sl in (sec1, sec2, sec3, sec4)]
content = [sl for sl in built if all(sl is not t for t in (sec1, sec2, sec3, sec4))]
print("content pages:", len(content))
assert len(content) == 14 and len(built) == 18, (len(content), len(built))

ordered = [cover, toc, sec1,
           content[0], content[1], content[2],
           sec2,
           content[3], content[4], content[5], content[6], content[7], content[8],
           sec3,
           content[9], content[10], content[11],
           sec4,
           content[12], content[13],
           end]
assert len(ordered) == TOTAL, (len(ordered), TOTAL)

drop = [t for t in tpl if all(t is not k for k in (cover, toc, end))]
set_order(prs, ordered, drop)

prs.save(OUT)
print("saved:", OUT)
print("slides:", len(Presentation(OUT).slides))
